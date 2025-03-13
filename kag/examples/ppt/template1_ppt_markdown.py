from kag.examples.ppt.template1.ppt_manager import PPTManager
from template1.ppt_utils import PPTUtils
import requests
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt


if __name__ == "__main__":
    # 创建一个新的 PowerPoint 演示文稿
    manager = PPTManager('template1/中国风背景.pptx')
    ppt_utils = PPTUtils()
    # 提取所有表格内容
    all_tables = ppt_utils.extract_markdown_tables('农林牧渔.md')

    for i, (table_content, merge_cells) in enumerate(all_tables):
        title = f"Table {i + 1}"  # 为每个表格自动生成标题
        # 添加表格到单独的幻灯片
        manager.creators['table'].add_table_slide_with_merge(
            title=title,
            content=None,
            table_content=table_content,
            merge_cells=merge_cells
        )

    # # 提取所有图片信息
    # images = ppt_utils.extract_markdown_images("农林牧渔.md")
    #
    # for idx, img in enumerate(images, 1):
    #     print(f"图片 {idx}:")
    #     print(f"标题：{img.get('title', '无标题')}")
    #     print(f"URL：{img['url']}")
    #     print("-" * 30)

    # # 示例数据
    # title = "相对沪深 300 表现图"
    # image_url = "https://cdn.noedgeai.com/019588ee-2816-7361-bede-ccca5c8ab74d_0.jpg?x=1102&y=1481&w=472&h=304&r=0"
    #
    # # 添加一个幻灯片并插入图片相关数据
    # manager.creators['image'].add_image_slide_online(title, image_url)
    #
    # 保存演示文稿
    manager.save("example_ppt_with_image_url.pptx")
